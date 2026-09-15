# -*- coding: utf-8 -*-
import io
from datetime import datetime
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


class ExcelExportService:
    """
    정기 업무 투입 현황 리포트 Excel 내보내기 서비스
    다중 시트 구성:
      1. 요약 현황
      2. 팀원별 투입 현황
      3. 고객사별 투입 현황
      4. 상세 작업 원장
    """

    @classmethod
    def generate_report(cls, df: pd.DataFrame, title_suffix: str = "") -> bytes:
        if df.empty:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "데이터 없음"
            ws["A1"] = "조회된 데이터가 없습니다."
            output = io.BytesIO()
            wb.save(output)
            output.seek(0)
            return output.getvalue()

        # 데이터 안전 복사 및 전처리
        data = df.copy()
        for dt_col in ["start_time", "end_time"]:
            if dt_col in data.columns:
                data[dt_col] = pd.to_datetime(data[dt_col], errors="coerce")

        wb = openpyxl.Workbook()
        # 기본 시트 제거를 위해 참조 보관
        default_sheet = wb.active

        # 색상 및 스타일 정의
        navy_dark = "002D42"
        navy_light = "F0F9FF"
        navy_sub = "0284C7"
        gray_bg = "F8FAFC"
        border_light = Side(border_style="thin", color="E2E8F0")
        box_border = Border(left=border_light, right=border_light, top=border_light, bottom=border_light)
        
        font_title = Font(name="맑은 고딕", size=15, bold=True, color="002D42")
        font_subtitle = Font(name="맑은 고딕", size=10, color="64748B")
        font_th = Font(name="맑은 고딕", size=10.5, bold=True, color="FFFFFF")
        font_td = Font(name="맑은 고딕", size=10)
        font_td_bold = Font(name="맑은 고딕", size=10, bold=True)
        
        fill_th = PatternFill(start_color=navy_dark, end_color=navy_dark, fill_type="solid")
        fill_sub_th = PatternFill(start_color=navy_sub, end_color=navy_sub, fill_type="solid")
        fill_zebra = PatternFill(start_color=gray_bg, end_color=gray_bg, fill_type="solid")
        fill_white = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

        # =========================================================
        # Sheet 1: 요약 현황
        # =========================================================
        ws1 = wb.create_sheet(title="요약 현황")
        ws1.views.sheetView[0].showGridLines = True
        
        now_str = datetime.now().strftime("%Y-%m-%d %H:%M")
        period_text = title_suffix if title_suffix else "전체 기간"
        
        ws1["A1"] = f"📊 정기 업무 투입 현황 리포트 ({period_text})"
        ws1["A1"].font = font_title
        ws1["A2"] = f"출력 일시: {now_str} | 대상 레코드: {len(data):,}건"
        ws1["A2"].font = font_subtitle

        # 핵심 지표 산출
        active_df = data[data["status"].isin(["COMPLETED", "PENDING"])] if "status" in data.columns else data
        tot_hours = round(active_df["actual_hours"].sum(), 1) if "actual_hours" in active_df.columns else 0.0
        tot_tasks = len(active_df)
        comp_tasks = len(active_df[active_df["status"] == "COMPLETED"]) if "status" in active_df.columns else 0
        pend_tasks = len(active_df[active_df["status"] == "PENDING"]) if "status" in active_df.columns else 0
        tot_workers = active_df["worker_name"].nunique() if "worker_name" in active_df.columns else 0
        tot_clients = active_df["client_name"].nunique() if "client_name" in active_df.columns else 0
        night_hours = round(active_df[active_df["is_night_work"] == True]["actual_hours"].sum(), 1) if "is_night_work" in active_df.columns else 0.0
        weekend_hours = round(active_df[active_df["is_weekend_work"] == True]["actual_hours"].sum(), 1) if "is_weekend_work" in active_df.columns else 0.0

        summary_metrics = [
            ("총 투입 공수", f"{tot_hours:,.1f} 시간"),
            ("총 지원 건수", f"{tot_tasks:,} 건 (완료: {comp_tasks:,} / 진행: {pend_tasks:,})"),
            ("투입 인원", f"{tot_workers} 명"),
            ("지원 고객사", f"{tot_clients} 개사"),
            ("야간 근무 공수", f"{night_hours:,.1f} 시간"),
            ("주말 근무 공수", f"{weekend_hours:,.1f} 시간"),
        ]

        ws1["A4"] = "구분"
        ws1["B4"] = "실적 지표"
        for col_letter in ["A", "B"]:
            cell = ws1[f"{col_letter}4"]
            cell.font = font_th
            cell.fill = fill_th
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = box_border

        for r_idx, (m_label, m_val) in enumerate(summary_metrics, start=5):
            c1 = ws1[f"A{r_idx}"]
            c2 = ws1[f"B{r_idx}"]
            c1.value = m_label
            c2.value = m_val
            c1.font = font_td_bold
            c2.font = font_td
            c1.border = box_border
            c2.border = box_border
            c1.fill = fill_zebra if r_idx % 2 == 1 else fill_white
            c2.fill = fill_zebra if r_idx % 2 == 1 else fill_white
            c1.alignment = Alignment(horizontal="left", vertical="center")
            c2.alignment = Alignment(horizontal="right", vertical="center")

        # 팀별 집계 표 (Sheet 1 하단)
        if "worker_team" in active_df.columns:
            start_r = 13
            ws1[f"A{start_r}"] = "소속팀별 공수 점유율"
            ws1[f"A{start_r}"].font = font_title
            start_r += 1

            t_cols = ["소속팀", "투입 인원(명)", "작업 건수", "총 투입공수(h)", "점유율(%)"]
            for c_i, c_name in enumerate(t_cols, start=1):
                c_cell = ws1.cell(row=start_r, column=c_i, value=c_name)
                c_cell.font = font_th
                c_cell.fill = fill_sub_th
                c_cell.alignment = Alignment(horizontal="center", vertical="center")
                c_cell.border = box_border

            team_grp = active_df.groupby("worker_team").agg(
                workers=("worker_name", "nunique"),
                tasks=("worker_name", "count"),
                hours=("actual_hours", "sum")
            ).reset_index().sort_values(by="hours", ascending=False)

            for t_idx, row in team_grp.iterrows():
                start_r += 1
                share = (row["hours"] / tot_hours * 100.0) if tot_hours > 0 else 0.0
                vals = [row["worker_team"], row["workers"], row["tasks"], round(row["hours"], 1), f"{share:.1f}%"]
                for c_i, v in enumerate(vals, start=1):
                    cell = ws1.cell(row=start_r, column=c_i, value=v)
                    cell.font = font_td
                    cell.border = box_border
                    cell.fill = fill_zebra if start_r % 2 == 1 else fill_white
                    cell.alignment = Alignment(horizontal="center" if c_i <= 3 else "right", vertical="center")

        # =========================================================
        # Sheet 2: 팀원별 투입 현황
        # =========================================================
        ws2 = wb.create_sheet(title="팀원별 투입 현황")
        ws2.views.sheetView[0].showGridLines = True
        
        ws2["A1"] = f"👥 팀원별 업무 투입 실적 ({period_text})"
        ws2["A1"].font = font_title

        w_headers = ["담당자", "소속팀", "직급", "총 작업건수", "총 투입공수(h)", "일반공수(h)", "야간공수(h)", "주말공수(h)", "완료건수", "진행건수"]
        for c_i, h_name in enumerate(w_headers, start=1):
            cell = ws2.cell(row=3, column=c_i, value=h_name)
            cell.font = font_th
            cell.fill = fill_th
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = box_border

        if "worker_name" in active_df.columns:
            w_rows = []
            for w_name, w_df in active_df.groupby("worker_name"):
                w_team = w_df["worker_team"].iloc[0] if "worker_team" in w_df.columns and not w_df["worker_team"].isna().all() else ""
                w_title = w_df["worker_title"].iloc[0] if "worker_title" in w_df.columns and not w_df["worker_title"].isna().all() else ""
                tot_h = round(w_df["actual_hours"].sum(), 1) if "actual_hours" in w_df.columns else 0.0
                night_h = round(w_df[w_df["is_night_work"] == True]["actual_hours"].sum(), 1) if "is_night_work" in w_df.columns else 0.0
                week_h = round(w_df[w_df["is_weekend_work"] == True]["actual_hours"].sum(), 1) if "is_weekend_work" in w_df.columns else 0.0
                day_h = max(0.0, round(tot_h - (night_h + week_h), 1))
                t_cnt = len(w_df)
                c_cnt = len(w_df[w_df["status"] == "COMPLETED"]) if "status" in w_df.columns else 0
                p_cnt = len(w_df[w_df["status"] == "PENDING"]) if "status" in w_df.columns else 0
                w_rows.append((w_name, w_team, w_title, t_cnt, tot_h, day_h, night_h, week_h, c_cnt, p_cnt))

            w_rows.sort(key=lambda x: x[4], reverse=True)

            for r_i, w_data in enumerate(w_rows, start=4):
                for c_i, v in enumerate(w_data, start=1):
                    cell = ws2.cell(row=r_i, column=c_i, value=v)
                    cell.font = font_td
                    cell.border = box_border
                    cell.fill = fill_zebra if r_i % 2 == 1 else fill_white
                    cell.alignment = Alignment(horizontal="center" if c_i <= 3 else "right", vertical="center")

        # =========================================================
        # Sheet 3: 고객사별 투입 현황
        # =========================================================
        ws3 = wb.create_sheet(title="고객사별 투입 현황")
        ws3.views.sheetView[0].showGridLines = True

        ws3["A1"] = f"🏢 고객사별 지원 실적 ({period_text})"
        ws3["A1"].font = font_title

        c_headers = ["고객사명", "총 작업건수", "총 투입공수(h)", "점유율(%)", "투입 인원수", "주요 담당자"]
        for c_i, h_name in enumerate(c_headers, start=1):
            cell = ws3.cell(row=3, column=c_i, value=h_name)
            cell.font = font_th
            cell.fill = fill_th
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = box_border

        if "client_name" in active_df.columns:
            c_rows = []
            for c_name, c_df in active_df.groupby("client_name"):
                tot_h = round(c_df["actual_hours"].sum(), 1) if "actual_hours" in c_df.columns else 0.0
                t_cnt = len(c_df)
                share = (tot_h / tot_hours * 100.0) if tot_hours > 0 else 0.0
                u_workers = [str(w) for w in c_df["worker_name"].dropna().unique() if str(w).strip()]
                w_count = len(u_workers)
                top_workers_str = ", ".join(u_workers[:4]) + (" 외" if len(u_workers) > 4 else "")
                c_rows.append((c_name, t_cnt, tot_h, share, w_count, top_workers_str))

            c_rows.sort(key=lambda x: x[2], reverse=True)

            for r_i, c_data in enumerate(c_rows, start=4):
                vals = [c_data[0], c_data[1], c_data[2], f"{c_data[3]:.1f}%", c_data[4], c_data[5]]
                for c_i, v in enumerate(vals, start=1):
                    cell = ws3.cell(row=r_i, column=c_i, value=v)
                    cell.font = font_td
                    cell.border = box_border
                    cell.fill = fill_zebra if r_i % 2 == 1 else fill_white
                    if c_i == 1:
                        cell.alignment = Alignment(horizontal="left", vertical="center")
                    elif c_i == 6:
                        cell.alignment = Alignment(horizontal="left", vertical="center")
                    else:
                        cell.alignment = Alignment(horizontal="right", vertical="center")

        # =========================================================
        # Sheet 4: 상세 작업 원장
        # =========================================================
        ws4 = wb.create_sheet(title="상세 작업 원장")
        ws4.views.sheetView[0].showGridLines = True

        ws4["A1"] = f"📋 상세 작업 원장 ({period_text})"
        ws4["A1"].font = font_title

        raw_headers = [
            "구분", "담당자", "소속팀", "직급", "고객사", "작업내용",
            "시작일시", "종료일시", "예정(h)", "실제소요(h)", "상태",
            "야간작업", "주말작업"
        ]
        for c_i, h_name in enumerate(raw_headers, start=1):
            cell = ws4.cell(row=3, column=c_i, value=h_name)
            cell.font = font_th
            cell.fill = fill_th
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = box_border

        disp_raw = data.sort_values(by="start_time", ascending=False) if "start_time" in data.columns else data
        for r_i, (_, r) in enumerate(disp_raw.iterrows(), start=4):
            st_str = r["start_time"].strftime("%Y-%m-%d %H:%M") if pd.notna(r.get("start_time")) and hasattr(r["start_time"], "strftime") else str(r.get("start_time", ""))[:16]
            ed_str = r["end_time"].strftime("%Y-%m-%d %H:%M") if pd.notna(r.get("end_time")) and hasattr(r["end_time"], "strftime") else str(r.get("end_time", ""))[:16]

            row_vals = [
                str(r.get("log_type", "작업")),
                str(r.get("worker_name", "")),
                str(r.get("worker_team", "")),
                str(r.get("worker_title", "")),
                str(r.get("client_name", "")),
                str(r.get("task_description", "")),
                st_str,
                ed_str,
                round(float(r.get("estimated_hours") or 0), 1),
                round(float(r.get("actual_hours") or 0), 1),
                "완료" if str(r.get("status")) == "COMPLETED" else "진행중",
                "O" if bool(r.get("is_night_work")) else "",
                "O" if bool(r.get("is_weekend_work")) else "",
            ]

            for c_i, v in enumerate(row_vals, start=1):
                cell = ws4.cell(row=r_i, column=c_i, value=v)
                cell.font = font_td
                cell.border = box_border
                cell.fill = fill_zebra if r_i % 2 == 1 else fill_white
                if c_i in [1, 2, 3, 4, 11, 12, 13]:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                elif c_i in [5, 6]:
                    cell.alignment = Alignment(horizontal="left", vertical="center")
                elif c_i in [7, 8]:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="right", vertical="center")

        # 기본 빈 시트 제거
        if default_sheet in wb.worksheets:
            wb.remove(default_sheet)

        # 모든 시트 자동 열 너비 계산 및 패딩 조정
        for ws in wb.worksheets:
            for col in ws.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    if cell.row < 3:
                        continue
                    val_str = str(cell.value or "")
                    # 한글 등 멀티바이트 글자 길이 고려
                    val_len = sum(2 if ord(ch) > 127 else 1 for ch in val_str)
                    if val_len > max_len:
                        max_len = val_len
                ws.column_dimensions[col_letter].width = max(max_len + 3, 11)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output.getvalue()
